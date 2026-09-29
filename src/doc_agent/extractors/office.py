"""Офисные форматы: pdf, docx, xlsx, csv, txt, md, pptx (опц.)."""
from __future__ import annotations

import contextlib
import csv
import warnings
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document as docx_document
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook

from doc_agent.config import load_settings
from doc_agent.extractors.base import ExtractResult, node_heading, node_image, node_paragraph, node_table
from doc_agent.extractors.registry import register


def _split_paragraphs(block: str) -> list[str]:
    parts = [p.strip() for p in block.replace("\r", "").split("\n")]
    return [p for p in parts if p]


@register("pdf")
def extract_pdf(path: Path, out_dir: Path) -> ExtractResult:
    tree: dict = {"version": "1.0", "meta": {"format": "pdf"}, "children": []}
    tables: list[dict] = []
    texts: list[str] = []
    page_spans: list[tuple[int, int]] = []
    char_pos = 0
    pages = 0
    with pdfplumber.open(str(path)) as pdf:
        pages = len(pdf.pages)
        tree["meta"]["pages"] = pages
        for pno, page in enumerate(pdf.pages, start=1):
            page_spans.append((char_pos, pno))
            block = page.extract_text() or ""
            for para in _split_paragraphs(block):
                node = node_paragraph(para, page=pno)
                tree["children"].append(node)
                texts.append(para)
                char_pos += len(para) + 2
            for tno, tbl in enumerate(page.extract_tables() or [], start=1):
                rows = [[(c or "").strip() for c in row] for row in tbl if row]
                if not rows:
                    continue
                name = f"Таблица {pno}.{tno}"
                tree["children"].append(node_table(name, rows, page=pno))
                tables.append({"name": name, "rows": rows})
    tree["meta"]["page_spans"] = page_spans
    return ExtractResult(text="\n\n".join(texts), tree=tree, tables=tables, pages=pages)


def _iter_docx_body(doc: Any):
    from docx.oxml.ns import qn

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, doc)


@register("docx")
def extract_docx(path: Path, out_dir: Path) -> ExtractResult:
    doc = docx_document(str(path))
    tree: dict = {"version": "1.0", "meta": {"format": "docx"}, "children": []}
    tables: list[dict] = []
    texts: list[str] = []
    tno = 0
    for item in _iter_docx_body(doc):
        if isinstance(item, DocxParagraph):
            text = item.text.strip()
            if not text:
                continue
            if item.style is not None and str(item.style.name or "").lower().startswith("heading"):
                level = 1
                try:
                    level = int(str(item.style.name).split()[-1])
                except ValueError:
                    pass
                tree["children"].append(node_heading(level, text))
            else:
                tree["children"].append(node_paragraph(text))
            texts.append(text)
        elif isinstance(item, DocxTable):
            tno += 1
            rows = [[(c.text or "").strip() for c in row.cells] for row in item.rows]
            name = f"Таблица {tno}"
            tree["children"].append(node_table(name, rows))
            tables.append({"name": name, "rows": rows})
    # Картинки из docx (word/media/*)
    import zipfile

    img_dir = out_dir
    img_dir.mkdir(parents=True, exist_ok=True)
    try:
        with zipfile.ZipFile(path) as zf:
            for i, name in enumerate(sorted(zf.namelist())):
                if name.startswith("word/media/"):
                    target = img_dir / f"img_{i + 1}{Path(name).suffix}"
                    target.write_bytes(zf.read(name))
                    tree["children"].append(node_image(str(target), caption=name))
    except Exception:  # noqa: BLE001
        pass
    return ExtractResult(text="\n\n".join(texts), tree=tree, tables=tables, pages=1)


@register("xlsx", "xlsm")
def extract_xlsx(path: Path, out_dir: Path) -> ExtractResult:
    # openpyxl сыплет UserWarning о print-областях/стилях («Print area cannot
    # be set…», «Workbook contains no default style…») — это шум, забивающий
    # консоль и логи. Глушится по флагу ingest.suppress_extract_warnings.
    suppress = load_settings().ingest.suppress_extract_warnings
    ctx = warnings.catch_warnings() if suppress else contextlib.nullcontext()
    with ctx:
        if suppress:
            warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(str(path), data_only=True, read_only=True)
    tree: dict = {"version": "1.0", "meta": {"format": "xlsx", "pages": 1}, "children": []}
    tables: list[dict] = []
    texts: list[str] = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        name = sheet.title or "Таблица"
        tree["children"].append(node_table(name, rows))
        tables.append({"name": name, "rows": rows})
    # v2.7: строки листов сериализуются в контекстном виде (contextual
    # retrieval): «Заголовок: значение; …» вместо слепого «;».join — пустые
    # ячейки больше не сдвигают колонки, у каждой строки виден смысл
    # колонок, маркер листа и шапка идут первыми строками листа.
    texts = []
    for _name, _header, _body in table_text_blocks(tables):
        block = f"=== Лист: {_name} ===\n{_body}"
        texts.append(block)
    wb.close()
    return ExtractResult(text="\n\n".join(texts), tree=tree, tables=tables, pages=1)


def _looks_numeric(value: str) -> bool:
    """Похоже ли значение на число/дату (для детекции строки шапки)."""
    v = value.strip().replace(" ", "").replace(",", ".")
    if not v:
        return False
    try:
        float(v)
        return True
    except ValueError:
        return any(ch.isdigit() for ch in v) and len(v) <= 12


def table_text_blocks(tables: list[dict]) -> list[tuple[str, str | None, str]]:
    """Таблицы -> [(имя_листа, шапка|None, текст_листа)] (v2.7).

    Публичный хелпер: используется и при извлечении (полный текст файла),
    и при табличном чанкинге в ingest (chunk_lines по каждому листу со
    своей шапкой). Шапка — первая из первых 10 строк, где >= 2 заполненных
    ячейки и все они текстовые, короткие и не числа; строки данных
    сериализуются как «Колонка: значение; …».
    """
    cfg = load_settings().chunking
    blocks: list[tuple[str, str | None, str]] = []
    for tbl in tables:
        name = str(tbl.get("name") or "Таблица")
        rows: list[list[str]] = [list(map(str, r)) for r in tbl.get("rows", [])]
        if not rows:
            continue
        header: str | None = None
        header_cells: list[str] = []
        data_start = 0
        if cfg.table_aware:
            for i, row in enumerate(rows[:10]):
                filled = [c for c in row if c]
                if len(filled) >= 2 and all(
                    not _looks_numeric(c) and len(c) <= 60 for c in filled
                ):
                    header = "; ".join(c for c in row if c)
                    header_cells = list(row)
                    data_start = i + 1
                    break
        lines: list[str] = []
        if header:
            lines.append(header)
        for row in rows[:data_start - 1] if data_start else []:
            lines.append("; ".join(c for c in row if c))
        for row in rows[data_start:]:
            if header:
                pairs = []
                for j, val in enumerate(row):
                    if not val:
                        continue
                    col = header_cells[j] if j < len(header_cells) and header_cells[j] \
                        else f"Колонка {j + 1}"
                    pairs.append(f"{col}: {val}")
                if pairs:
                    lines.append("; ".join(pairs))
            else:
                lines.append("; ".join(c for c in row if c))
        blocks.append((name, header, "\n".join(ln for ln in lines if ln)))
    return blocks


@register("csv")
def extract_csv(path: Path, out_dir: Path) -> ExtractResult:
    rows: list[list[str]] = []
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            with open(path, encoding=encoding, newline="") as f:
                rows = [[cell.strip() for cell in row] for row in csv.reader(f) if row]
            break
        except UnicodeDecodeError:
            continue
    tree: dict = {"version": "1.0", "meta": {"format": "csv", "pages": 1}, "children": []}
    if rows:
        tree["children"].append(node_table(path.stem, rows))
    text = "\n".join("; ".join(r) for r in rows)
    return ExtractResult(text=text, tree=tree, tables=[{"name": path.stem, "rows": rows}], pages=1)


@register("txt", "md")
def extract_text(path: Path, out_dir: Path) -> ExtractResult:
    raw = ""
    for encoding in ("utf-8", "cp1251"):
        try:
            raw = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    tree: dict = {"version": "1.0", "meta": {"format": path.suffix.lstrip(".") or "txt", "pages": 1}, "children": []}
    paras = _split_paragraphs(raw)
    for para in paras:
        if para.startswith("#"):
            tree["children"].append(node_heading(1, para.lstrip("# ")))
        else:
            tree["children"].append(node_paragraph(para))
    return ExtractResult(text="\n\n".join(paras), tree=tree, pages=1)


@register("pptx")
def extract_pptx(path: Path, out_dir: Path) -> ExtractResult:
    try:
        from pptx import Presentation  # type: ignore[import-untyped]
    except ImportError:
        raise RuntimeError("python-pptx не установлен (пакет pptx)")  # noqa: B904
    prs = Presentation(str(path))
    texts: list[str] = []
    tree: dict = {"version": "1.0", "meta": {"format": "pptx", "pages": len(prs.slides)}, "children": []}
    for sno, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        tree["children"].append(node_paragraph(t, page=sno))
                        texts.append(t)
    return ExtractResult(text="\n\n".join(texts), tree=tree, pages=len(prs.slides))
