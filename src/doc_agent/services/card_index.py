"""Картотека: выборка с фильтрами + экспорт XLSX (15 колонок, автофильтр) — DD-3."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from doc_agent.db.models import File

# Ровно 15 колонок (DD-3)
CARD_COLUMNS: list[str] = [
    "ID", "Имя файла", "Формат", "Город", "Заказчик", "Договор №", "Договор дата",
    "Тип документа", "Номер документа", "Дата документа", "Сумма", "Статус",
    "OCR уверенность", "Описание", "Источник",
]

_FIELD_BY_COLUMN = [
    "id", "original_name", "format", "city", "customer", "contract_no", "contract_date",
    "doc_type", "doc_number", "doc_date", "amount", "status",
    "ocr_confidence", "description", "source",
]


def query_card(db: Session, *, city: str | None = None, customer: str | None = None,
               contract_no: str | None = None, doc_type: str | None = None,
               date_from: str | None = None, date_to: str | None = None,
               limit: int = 50, offset: int = 0) -> dict[str, Any]:
    stmt = select(File)
    if city:
        stmt = stmt.where(File.city == city)
    if customer:
        stmt = stmt.where(File.customer.ilike(f"%{customer}%"))
    if contract_no:
        stmt = stmt.where(File.contract_no == contract_no)
    if doc_type:
        stmt = stmt.where(File.doc_type == doc_type)
    if date_from:
        stmt = stmt.where(File.doc_date >= date_from)
    if date_to:
        stmt = stmt.where(File.doc_date <= date_to)
    total = db.execute(select(func.count()).select_from(stmt.subquery())).scalar_one()
    rows = db.execute(stmt.order_by(File.id).limit(limit).offset(offset)).scalars().all()
    items = [
        {k: getattr(f, k) for k in _FIELD_BY_COLUMN}
        for f in rows
    ]
    return {"total": total, "items": items}


def _row_values(item: dict) -> list[Any]:
    values: list[Any] = []
    for field_name in _FIELD_BY_COLUMN:
        v = item.get(field_name)
        if field_name == "ocr_confidence" and isinstance(v, float):
            v = round(v, 3)
        values.append(v)
    return values


def export_xlsx(items: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Картотека"
    ws.append(CARD_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for item in items:
        ws.append(_row_values(item))
    # Автофильтр по всей таблице (DD-3)
    last_row = max(ws.max_row, 1)
    ws.auto_filter.ref = f"A1:O{last_row}"
    widths = [6, 36, 8, 16, 26, 14, 14, 12, 16, 14, 14, 12, 12, 42, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
