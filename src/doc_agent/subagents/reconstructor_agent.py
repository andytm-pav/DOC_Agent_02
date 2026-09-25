"""Reconstructor subagent for document reproduction."""

import json
from pathlib import Path

import structlog

logger = structlog.get_logger()


class ReconstructorAgent:
    """Reconstructs documents in different formats."""

    async def run(self, tree_path: str, target_format: str, output_path: str) -> str:
        """Reconstruct a document from its tree."""
        with open(tree_path, encoding="utf-8") as f:
            tree = json.load(f)

        if target_format == "docx":
            return self._to_docx(tree, output_path)
        elif target_format == "xlsx":
            return self._to_xlsx(tree, output_path)
        elif target_format == "md":
            return self._to_md(tree, output_path)
        else:
            raise ValueError(f"Unsupported format: {target_format}")

    def _to_docx(self, tree: dict, output_path: str) -> str:
        """Convert tree to DOCX."""
        from docx import Document

        doc = Document()
        for node in tree.get("children", []):
            if node["type"] == "heading":
                doc.add_heading(node.get("text", ""), level=node.get("level", 1))
            elif node["type"] == "paragraph":
                doc.add_paragraph(node.get("text", ""))
            elif node["type"] == "table" and node.get("table"):
                table_data = node["table"]
                rows = table_data.get("rows", [])
                if rows:
                    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    for i, row in enumerate(rows):
                        for j, cell in enumerate(row):
                            table.cell(i, j).text = str(cell)

        doc.save(output_path)
        return output_path

    def _to_xlsx(self, tree: dict, output_path: str) -> str:
        """Convert tree to XLSX."""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Документ"

        row_idx = 1
        for node in tree.get("children", []):
            if node["type"] == "paragraph":
                ws.cell(row=row_idx, column=1, value=node.get("text", ""))
                row_idx += 1
            elif node["type"] == "table" and node.get("table"):
                table_data = node["table"]
                for row in table_data.get("rows", []):
                    for col_idx, cell in enumerate(row, 1):
                        ws.cell(row=row_idx, column=col_idx, value=str(cell))
                    row_idx += 1

        wb.save(output_path)
        return output_path

    def _to_md(self, tree: dict, output_path: str) -> str:
        """Convert tree to Markdown."""
        lines = []
        for node in tree.get("children", []):
            if node["type"] == "heading":
                level = node.get("level", 1)
                lines.append(f"{'#' * level} {node.get('text', '')}")
            elif node["type"] == "paragraph":
                lines.append(node.get("text", ""))
            elif node["type"] == "table" and node.get("table"):
                rows = node["table"].get("rows", [])
                if rows:
                    # Header
                    lines.append("| " + " | ".join(rows[0]) + " |")
                    lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
                    for row in rows[1:]:
                        lines.append("| " + " | ".join(row) + " |")
            lines.append("")

        Path(output_path).write_text("\n".join(lines), encoding="utf-8")
        return output_path
