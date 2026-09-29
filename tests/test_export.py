"""test_export.py: /card/export.xlsx — валидный XLSX (13.2, DD-3)."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from conftest import upload_and_wait


def test_export_valid_xlsx(client, user_token: str, sample_files: dict) -> None:
    upload_and_wait(client, user_token, sample_files["pdf"])
    r = client.get(
        "/card/export.xlsx",
        headers={"Authorization": f"Bearer {user_token}"},
    )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert len(headers) == 15
    assert ws.auto_filter.ref is not None
    assert ws.max_row >= 2
