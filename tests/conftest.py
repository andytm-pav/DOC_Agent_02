"""Фикстуры (13.1): client, tmp_data_dir, admin_token/user_token, sample_files.

Конфиг тестов: llm.provider=none — детерминированный офлайн-режим
(правила rules.yaml + экстрактивный репортер). LLM-пути тестируются в
test_offline.py (недоступный Ollama -> fallback).
"""
from __future__ import annotations

import os
import shutil
import time
from collections.abc import Iterator
from pathlib import Path

import pytest
import yaml

REPO_ROOT = Path(__file__).resolve().parents[1]
sys_path = REPO_ROOT / "src"


def _make_base_config(data_dir: Path, llm_provider: str = "none",
                      llm_base_url: str = "http://localhost:9") -> Path:
    config = {
        "app": {"host": "127.0.0.1", "port": 8000, "secret_key": "test-secret",
                "data_dir": str(data_dir), "admin_username": "admin", "admin_password": "admin"},
        "storage": {"local_paths": [str(data_dir / "storage")], "network_paths": []},
        "watchdog": {"enabled": False, "paths": [], "interval_sec": 30},
        "llm": {"provider": llm_provider, "base_url": llm_base_url, "api_key": "",
                "model_main": "qwen2.5:1.5b-instruct-q4_K_M", "model_light": "qwen2.5:1.5b-instruct-q4_K_M",
                "timeout_sec": 5, "temperature": 0.0},
        "embeddings": {"backend": "hashing", "model": "intfloat/multilingual-e5-small", "batch_size": 8},
        "vector": {"backend": "chroma", "collection": "chunks", "metric": "cosine"},
        "reranker": {"backend": "simple", "model": "cross-encoder/ms-marco-MiniLM-L-6-v2", "top_k": 5},
        "chunking": {"size_tokens": 128, "overlap_tokens": 16},
        "ocr": {"engine": "tesseract", "tesseract_cmd": "", "tesseract_langs": ["rus", "eng"],
                "min_confidence": 0.80, "reject_threshold": 0.50, "process_embedded_images": False},
        "groundrails": {"enabled": True, "cut_ratio": 0.30},
        "a2a": {"enabled": True, "agent_name": "document-agent", "auth_required": False},
        "limits": {"max_file_mb": 100, "max_query_len": 2000},
        "auth": {"token_expire_minutes": 60,
                 "api_keys": [{"key": "test-api-key-user", "name": "external-system", "role": "user"},
                              {"key": "test-api-key-admin", "name": "external-admin", "role": "admin"}]},
    }
    cfg_dir = data_dir.parent / "config"
    cfg_dir.mkdir(parents=True, exist_ok=True)
    cfg_path = cfg_dir / "config.yaml"
    cfg_path.write_text(yaml.safe_dump(config, allow_unicode=True), encoding="utf-8")
    return cfg_path


@pytest.fixture()
def tmp_data_dir(tmp_path: Path) -> Iterator[Path]:
    data_dir = tmp_path / "data"
    for sub in ("storage", "sqlite", "chroma", "logs", "tmp", "inbox"):
        (data_dir / sub).mkdir(parents=True, exist_ok=True)
    cfg_path = _make_base_config(data_dir)
    os.environ["DOC_AGENT_CONFIG"] = str(cfg_path)
    # сброс кэшей движков/векторного store между тестами
    from doc_agent.db import session as db_session
    from doc_agent.vector import chroma_client

    db_session._engines.clear()
    db_session._factories.clear()
    chroma_client.reset_vector_store()
    yield data_dir
    os.environ.pop("DOC_AGENT_CONFIG", None)
    db_session._engines.clear()
    db_session._factories.clear()
    chroma_client.reset_vector_store()
    shutil.rmtree(tmp_path, ignore_errors=True)


@pytest.fixture()
def client(tmp_data_dir: Path):
    from fastapi.testclient import TestClient

    from doc_agent.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture()
def admin_token(client) -> str:
    r = client.post("/auth/login", json={"username": "admin", "password": "admin"})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture()
def user_token(client, admin_token) -> str:
    # отдельный пользователь не обязателен: admin проходит все user-эндпоинты,
    # но для чистоты создаём user через прямую запись.
    from doc_agent.db.migrations import hash_password
    from doc_agent.db.models import User
    from doc_agent.db.session import get_session_factory

    factory = get_session_factory()
    with factory() as db:
        if db.query(User).filter(User.username == "user").one_or_none() is None:
            db.add(User(username="user", password_hash=hash_password("user"), role="user"))
            db.commit()
    r = client.post("/auth/login", json={"username": "user", "password": "user"})
    return r.json()["access_token"]


# --- Генерация sample-файлов ---

KS2_TEXT = (
    "АКТ о приёмке выполненных работ КС-2 № 5 от 12.03.2024\n"
    "Город: Самара\n"
    "Заказчик: ООО СтройЗаказчик\n"
    "Договор № 123 от 01.02.2024\n"
    "Итого: 1 234 567,89 руб\n"
)


@pytest.fixture()
def sample_files(tmp_path: Path) -> dict[str, Path]:
    out = tmp_path / "samples"
    out.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}

    # PDF (reportlab — уже в зависимостях). Нужен шрифт с кириллицей:
    # системные TTF ищем кросс-платформенно (Linux/Windows).
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    _font_candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/times.ttf",
        "C:/Windows/Fonts/calibri.ttf",
        "C:/Windows/Fonts/DejaVuSans.ttf",
    ]
    font_name = "Helvetica"
    for cand in _font_candidates:
        if Path(cand).exists():
            pdfmetrics.registerFont(TTFont("FixtureSans", cand))
            font_name = "FixtureSans"
            break

    pdf_path = out / "ks2.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4)
    styles = getSampleStyleSheet()
    h2 = styles["Heading2"].clone("h2ru", fontName=font_name)
    body = styles["BodyText"].clone("bodyru", fontName=font_name)
    doc.build([
        Paragraph("АКТ о приёмке выполненных работ КС-2 № 5 от 12.03.2024", h2),
        Paragraph("Город: Самара", body),
        Paragraph("Заказчик: ООО СтройЗаказчик", body),
        Paragraph("Договор № 123 от 01.02.2024", body),
        Paragraph("Итого: 1 234 567,89 руб", body),
    ])
    files["pdf"] = pdf_path

    # DOCX с таблицей 3x3
    from docx import Document

    docx_path = out / "act_table.docx"
    d = Document()
    d.add_heading("Акт о приёмке выполненных работ КС-2 № 7", level=1)
    d.add_paragraph("Договор № 123. Итого: 50000,00 руб")
    table = d.add_table(rows=3, cols=3)
    data = [["Позиция", "Кол-во", "Сумма"], ["Работы", "10", "10000"], ["Материалы", "5", "40000"]]
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val
    d.save(str(docx_path))
    files["docx"] = docx_path

    # XLSX
    from openpyxl import Workbook

    xlsx_path = out / "report.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.append(["Наименование", "Сумма"])
    ws.append(["Работы КС-3 по договору 123", "70000"])
    wb.save(str(xlsx_path))
    files["xlsx"] = xlsx_path

    # JPG (маленькое изображение с текстом)
    try:
        from PIL import Image, ImageDraw

        img_path = out / "doc.jpg"
        img = Image.new("RGB", (400, 120), "white")
        drawer = ImageDraw.Draw(img)
        drawer.text((10, 50), "TTH 2024", fill="black")
        img.save(str(img_path))
        files["jpg"] = img_path
    except Exception:  # noqa: BLE001
        pass

    # EML
    from email.message import EmailMessage

    eml_path = out / "letter.eml"
    msg = EmailMessage()
    msg["Subject"] = "Счёт-фактура № 15"
    msg["From"] = "buh@customer.ru"
    msg["To"] = "agent@example.local"
    msg.set_content("Счёт-фактура № 15 от 05.04.2024 по договору № 123. Сумма 90000 руб.")
    eml_path.write_bytes(msg.as_bytes())
    files["eml"] = eml_path
    return files


def upload_and_wait(client, token: str, file_path: Path, timeout: float = 40.0) -> dict:
    """Загружает файл и ждёт завершения обработки. Возвращает данные GET /files/{id}."""
    with open(file_path, "rb") as fh:
        r = client.post(
            "/files/upload",
            files={"file": (file_path.name, fh)},
            headers={"Authorization": f"Bearer {token}"},
        )
    assert r.status_code == 200, r.text
    file_id = r.json()["file_id"]
    deadline = time.time() + timeout
    while time.time() < deadline:
        r2 = client.get(f"/files/{file_id}", headers={"Authorization": f"Bearer {token}"})
        data = r2.json()
        if data["file"]["status"] in {"done", "failed", "description_only"}:
            return data
        time.sleep(0.3)
    raise TimeoutError(f"Файл {file_path.name} не обработан за {timeout} с")
