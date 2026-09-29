"""test_card_index.py: 15 колонок в XLSX, автофильтр включён (13.2)."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from doc_agent.services.card_index import CARD_COLUMNS, export_xlsx


def test_columns_are_15() -> None:
    assert len(CARD_COLUMNS) == 15


def test_xlsx_autofilter() -> None:
    items = [{
        "id": 1, "original_name": "ks2.pdf", "format": "pdf", "city": "Самара",
        "customer": "ООО СтройЗаказчик", "contract_no": "123", "contract_date": "2024-02-01",
        "doc_type": "ks2", "doc_number": "5", "doc_date": "2024-03-12", "amount": 1234567.89,
        "status": "done", "ocr_confidence": None,
        "description": "Тип: ks2", "source": "user",
    }]
    content = export_xlsx(items)
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == CARD_COLUMNS
    assert len(headers) == 15
    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith("A1:O")
    assert ws.max_row == 2


def test_card_endpoint(client, user_token: str, sample_files: dict) -> None:
    from conftest import upload_and_wait

    upload_and_wait(client, user_token, sample_files["pdf"])
    r = client.get("/card", headers={"Authorization": f"Bearer {user_token}"})
    assert r.status_code == 200
    data = r.json()
    assert data["total"] >= 1
    assert len(data["items"][0]) == 15
