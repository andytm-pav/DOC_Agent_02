"""test_reconstruct.py: DOCX с таблицей 3x3 -> XLSX с 3x3 (13.2, DD-7)."""
from __future__ import annotations

import io
import time

from openpyxl import load_workbook

from conftest import upload_and_wait


def _wait_job_done(client, token: str, job_id: int, timeout: float = 20.0) -> None:
    deadline = time.time() + timeout
    while time.time() < deadline:
        r = client.get("/admin/jobs?limit=200", headers={"Authorization": f"Bearer {token}"})
        for job in r.json()["jobs"]:
            if job["id"] == job_id and job["status"] in {"done", "failed"}:
                assert job["status"] == "done", job["last_error"]
                return
        time.sleep(0.3)
    raise TimeoutError("reproduce job не завершился")


def test_docx_table_to_xlsx(client, admin_token: str, user_token: str, sample_files: dict) -> None:
    data = upload_and_wait(client, user_token, sample_files["docx"])
    file_id = data["file"]["id"]

    r = client.post(
        f"/files/{file_id}/reproduce",
        json={"target_format": "xlsx"},
        headers={"Authorization": f"Bearer {user_token}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert "download_url" in body
    job_id = int(body["download_url"].rstrip("/").split("/")[-1])
    _wait_job_done(client, admin_token, job_id)

    dl = client.get(
        f"/files/{file_id}/download/{job_id}",
        headers={"Authorization": f"Bearer {user_token}"},
    )
    assert dl.status_code == 200, dl.text
    wb = load_workbook(io.BytesIO(dl.content))
    ws = wb[wb.sheetnames[0]]
    assert ws.max_row == 3
    assert ws.max_column == 3
    assert ws.cell(row=1, column=1).value == "Позиция"


def test_reproduce_unsupported_format(client, user_token: str, sample_files: dict) -> None:
    data = upload_and_wait(client, user_token, sample_files["docx"])
    file_id = data["file"]["id"]
    r = client.post(
        f"/files/{file_id}/reproduce",
        json={"target_format": "exe"},
        headers={"Authorization": f"Bearer {user_token}"},
    )
    assert r.status_code == 422


def test_reproduce_md(client, user_token: str, sample_files: dict) -> None:
    data = upload_and_wait(client, user_token, sample_files["docx"])
    file_id = data["file"]["id"]
    r = client.post(
        f"/files/{file_id}/reproduce",
        json={"target_format": "md"},
        headers={"Authorization": f"Bearer {user_token}"},
    )
    assert r.status_code == 200
    job_id = int(r.json()["download_url"].rstrip("/").split("/")[-1])
    deadline = time.time() + 15
    status_code = 0
    while time.time() < deadline:
        dl = client.get(f"/files/{file_id}/download/{job_id}", headers={"Authorization": f"Bearer {user_token}"})
        status_code = dl.status_code
        if status_code == 200:
            assert "Акт" in dl.text or "# Акт" in dl.text
            return
        time.sleep(0.3)
    raise AssertionError(f"md не воспроизвёлся (status {status_code})")
